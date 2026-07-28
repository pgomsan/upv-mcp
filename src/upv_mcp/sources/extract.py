"""Extraccion de texto de los ficheros de PoliformaT.

Convierte un PDF, un PowerPoint o un Word en texto plano para que el modelo pueda
responder sobre su contenido. No es OCR: un PDF escaneado (paginas que en realidad
son imagenes) no tiene texto que extraer, y eso se dice en vez de devolver vacio.

INVARIANTE: este modulo no importa nada de `mcp`. Solo parseo.
"""

from __future__ import annotations

import io
from typing import Final

#: Tope de texto devuelto. Un tema de 40 transparencias son ~4 paginas; 40k
#: caracteres cubre documentos largos sin arrasar el contexto del cliente.
MAX_CHARS: Final = 40_000

#: Tipos que sabemos convertir. El resto se rechaza con un motivo claro.
PDF: Final = "application/pdf"
PPTX: Final = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
DOCX: Final = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

_TEXTO_PLANO: Final = frozenset(
    {"text/plain", "text/markdown", "text/csv", "text/html", "application/json"}
)

SOPORTADOS: Final = frozenset({PDF, PPTX, DOCX}) | _TEXTO_PLANO


class ExtractionError(RuntimeError):
    """No se pudo sacar texto del fichero. El mensaje explica por que."""


def _pdf(datos: bytes) -> str:
    from pypdf import PdfReader

    try:
        lector = PdfReader(io.BytesIO(datos))
    except Exception as exc:
        raise ExtractionError(f"El PDF no se pudo abrir: {exc}") from exc

    if lector.is_encrypted:
        raise ExtractionError("El PDF esta protegido con contrasena.")

    paginas = []
    for numero, pagina in enumerate(lector.pages, start=1):
        try:
            texto = (pagina.extract_text() or "").strip()
        except Exception:
            continue
        if texto:
            paginas.append(f"--- Pagina {numero} ---\n{texto}")

    if not paginas:
        raise ExtractionError(
            f"El PDF tiene {len(lector.pages)} paginas pero ninguna contiene texto "
            "extraible. Suele pasar con documentos escaneados o hechos de imagenes; "
            "haria falta OCR, que este servidor no hace."
        )
    return "\n\n".join(paginas)


def _pptx(datos: bytes) -> str:
    from pptx import Presentation

    try:
        presentacion = Presentation(io.BytesIO(datos))
    except Exception as exc:
        raise ExtractionError(f"La presentacion no se pudo abrir: {exc}") from exc

    diapositivas = []
    for numero, diapositiva in enumerate(presentacion.slides, start=1):
        trozos = [
            forma.text_frame.text.strip()
            for forma in diapositiva.shapes
            if forma.has_text_frame and forma.text_frame.text.strip()
        ]
        if trozos:
            diapositivas.append(f"--- Diapositiva {numero} ---\n" + "\n".join(trozos))

    if not diapositivas:
        raise ExtractionError(
            "La presentacion no tiene texto extraible: sus diapositivas son probablemente imagenes."
        )
    return "\n\n".join(diapositivas)


def _docx(datos: bytes) -> str:
    from docx import Document

    try:
        documento = Document(io.BytesIO(datos))
    except Exception as exc:
        raise ExtractionError(f"El documento no se pudo abrir: {exc}") from exc

    parrafos = [p.text.strip() for p in documento.paragraphs if p.text.strip()]
    for tabla in documento.tables:
        for fila in tabla.rows:
            celdas = [c.text.strip() for c in fila.cells if c.text.strip()]
            if celdas:
                parrafos.append(" | ".join(celdas))

    if not parrafos:
        raise ExtractionError("El documento esta vacio o no tiene texto extraible.")
    return "\n\n".join(parrafos)


def _plano(datos: bytes) -> str:
    texto = datos.decode("utf-8", errors="replace").strip()
    if not texto:
        raise ExtractionError("El fichero esta vacio.")
    return texto


def _formato_real(datos: bytes, content_type: str | None, nombre: str) -> str | None:
    """Decide el formato mirando primero los BYTES, no el nombre.

    En PoliformaT hay ficheros mal nombrados: uno llamado "PresentacionRIN.pdf" es
    en realidad un PowerPoint. Fiarse de la extension hacia que pypdf reventara con
    un error incomprensible, asi que manda la firma del fichero y solo despues el
    content-type y la extension.
    """
    tipo = (content_type or "").split(";")[0].strip().lower()
    minusculas = nombre.lower()

    if datos.startswith(b"%PDF"):
        return PDF
    if datos.startswith(b"PK\x03\x04"):
        # Los formatos de Office son zips. Para saber cual, se mira dentro.
        cabeza = datos[:4096]
        if b"ppt/" in cabeza or tipo == PPTX or minusculas.endswith(".pptx"):
            return PPTX
        if b"word/" in cabeza or tipo == DOCX or minusculas.endswith(".docx"):
            return DOCX
        return None  # Un zip normal: no hay texto que sacar.

    if tipo in (PDF, PPTX, DOCX):
        return tipo
    if tipo in _TEXTO_PLANO:
        return "texto"
    if minusculas.endswith(".pdf"):
        return PDF
    if minusculas.endswith((".txt", ".md", ".csv", ".json")):
        return "texto"
    return None


def extract(datos: bytes, content_type: str | None, nombre: str = "") -> tuple[str, bool]:
    """Devuelve `(texto, recortado)` a partir del contenido de un fichero.

    Un formato no soportado (ZIP, imagen, video) da un error explicando cual es,
    para que la respuesta pueda decir "esto no lo puedo leer" en vez de fallar sin
    mas.
    """
    formato = _formato_real(datos, content_type, nombre)

    if formato == PDF:
        texto = _pdf(datos)
    elif formato == PPTX:
        texto = _pptx(datos)
    elif formato == DOCX:
        texto = _docx(datos)
    elif formato == "texto":
        texto = _plano(datos)
    else:
        declarado = (content_type or "").split(";")[0].strip() or "desconocido"
        raise ExtractionError(
            f"No se puede extraer texto de este fichero (tipo '{declarado}'). "
            "Se soportan PDF, PowerPoint (.pptx), Word (.docx) y texto plano. "
            "Los ZIP, imagenes y videos hay que abrirlos desde su URL."
        )

    texto = "\n".join(linea.rstrip() for linea in texto.splitlines())
    if len(texto) > MAX_CHARS:
        return texto[:MAX_CHARS].rsplit(" ", 1)[0], True
    return texto, False
